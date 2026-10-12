"""Multi-format document parsers for PDF, DOCX, XLSX, CSV, and image files."""
import os
import hashlib
from typing import List, Dict, Any, Tuple
from pathlib import Path

# Parsers
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

import pypdf
import docx
import openpyxl
import pandas as pd

def calculate_file_hash(file_bytes: bytes) -> str:
    """Calculates SHA-256 hash of file content to detect duplicates and ensure integrity."""
    return hashlib.sha256(file_bytes).hexdigest()

def parse_pdf(file_path: str) -> List[Dict[str, Any]]:
    """Extracts pages and text from PDF using PyMuPDF if available, or pypdf as fallback."""
    pages_data = []
    if HAS_PYMUPDF:
        doc = fitz.open(file_path)
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text").strip()
            # Also extract simple tables if present
            tables_text = ""
            try:
                tables = page.find_tables()
                for table in tables:
                    df = table.to_pandas()
                    tables_text += "\n[Table: " + df.to_string(index=False) + "]\n"
            except Exception:
                pass
            
            full_content = text + ("\n" + tables_text if tables_text else "")
            if full_content.strip():
                pages_data.append({
                    "page_number": page_num + 1,
                    "section_title": f"Page {page_num + 1}",
                    "content": full_content
                })
        doc.close()
    else:
        reader = pypdf.PdfReader(file_path)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                pages_data.append({
                    "page_number": page_num + 1,
                    "section_title": f"Page {page_num + 1}",
                    "content": text.strip()
                })
    return pages_data

def parse_docx(file_path: str) -> List[Dict[str, Any]]:
    """Extracts structured text and tables from Word documents (.docx)."""
    doc = docx.Document(file_path)
    sections_data = []
    current_heading = "General"
    current_content = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            if current_content:
                sections_data.append({
                    "page_number": 1,
                    "section_title": current_heading,
                    "content": "\n".join(current_content)
                })
                current_content = []
            current_heading = text
        else:
            current_content.append(text)
            
    if current_content:
        sections_data.append({
            "page_number": 1,
            "section_title": current_heading,
            "content": "\n".join(current_content)
        })
        
    # Extract tables
    for t_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            table_rows.append([cell.text.strip() for cell in row.cells])
        if table_rows:
            headers = table_rows[0]
            rows_str = []
            for r in table_rows[1:]:
                row_repr = ", ".join([f"{h}: {val}" for h, val in zip(headers, r) if val])
                if row_repr:
                    rows_str.append(row_repr)
            table_text = f"[Specification Table {t_idx + 1}]:\n" + "\n".join(rows_str)
            sections_data.append({
                "page_number": 1,
                "section_title": f"Table {t_idx + 1}",
                "content": table_text
            })
            
    return sections_data

def parse_excel(file_path: str) -> List[Dict[str, Any]]:
    """Extracts tabular specification data from Excel sheets (.xlsx)."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheets_data = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        data = list(sheet.iter_rows(values_only=True))
        if not data:
            continue
            
        header = [str(col).strip() if col is not None else f"Col_{i}" for i, col in enumerate(data[0])]
        content_lines = []
        for row_idx, row in enumerate(data[1:], start=2):
            row_items = [f"{h}: {val}" for h, val in zip(header, row) if val is not None and str(val).strip()]
            if row_items:
                content_lines.append(f"Row {row_idx} -> " + " | ".join(row_items))
                
        if content_lines:
            sheets_data.append({
                "page_number": 1,
                "section_title": f"Sheet: {sheet_name}",
                "content": f"[Sheet: {sheet_name} Specification Table]\n" + "\n".join(content_lines)
            })
    return sheets_data

def parse_csv(file_path: str) -> List[Dict[str, Any]]:
    """Extracts specifications from CSV files."""
    df = pd.read_csv(file_path)
    lines = []
    headers = list(df.columns)
    for idx, row in df.iterrows():
        items = [f"{h}: {row[h]}" for h in headers if pd.notna(row[h])]
        if items:
            lines.append(f"Item {idx + 1} -> " + " | ".join(items))
    return [{
        "page_number": 1,
        "section_title": "CSV Specifications",
        "content": "[CSV Specification Table]\n" + "\n".join(lines)
    }]

def parse_document(file_path: str, file_type: str) -> List[Dict[str, Any]]:
    """Unified entrypoint for document parsing based on file extension."""
    ext = file_type.lower()
    if not ext.startswith("."):
        ext = "." + ext
        
    if ext == ".pdf":
        return parse_pdf(file_path)
    elif ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".xlsx":
        return parse_excel(file_path)
    elif ext == ".csv":
        return parse_csv(file_path)
    elif ext in [".txt", ".md"]:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read()
        return [{"page_number": 1, "section_title": "Document Content", "content": text}]
    elif ext in [".jpg", ".jpeg", ".png"]:
        # Fallback for images when OCR engine is optional
        return [{
            "page_number": 1,
            "section_title": "Image Asset",
            "content": f"[Visual Inspection Document: {Path(file_path).name}] (Photographic / visual evidence recorded for complaint or inspection file)"
        }]
    else:
        raise ValueError(f"Unsupported file format: {ext}")
