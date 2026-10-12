"""Excel and CSV report generation service."""
import io
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Dict, Any
from src.database.connection import get_db

def style_excel_worksheet(ws):
    """Applies clean enterprise styling to an openpyxl worksheet."""
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    border_side = Side(style='thin', color='CBD5E1')
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Style header row
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border
        
    # Style data rows and auto-adjust widths
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.font = data_font
            cell.border = cell_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

def generate_management_dashboard_excel() -> bytes:
    """Generates an executive Excel workbook with multiple tabs: Complaints, Evaluations, CAPAs, Products."""
    with get_db() as conn:
        complaints_df = pd.read_sql_query(
            """
            SELECT c.complaint_id, c.date_received, c.reporter_type, cust.name as customer_name,
                   p.sku, p.name as product_name, c.category, c.severity, c.urgency,
                   c.quantity_affected, c.current_status, u.full_name as assigned_owner
            FROM complaints c
            JOIN products p ON c.product_id = p.id
            LEFT JOIN customers cust ON c.customer_id = cust.id
            LEFT JOIN users u ON c.assigned_owner_id = u.id
            ORDER BY c.created_at DESC
            """,
            conn
        )
        
        evaluations_df = pd.read_sql_query(
            """
            SELECT e.evaluation_id, c.complaint_id, p.sku, p.name as product_name,
                   e.quality_score, e.scoring_status, e.mandatory_failures_count,
                   e.evidence_completeness_pct, e.manager_approval_status, e.created_at
            FROM evaluations e
            JOIN complaints c ON e.complaint_id = c.id
            JOIN products p ON e.product_id = p.id
            ORDER BY e.created_at DESC
            """,
            conn
        )
        
        capas_df = pd.read_sql_query(
            """
            SELECT cp.id as capa_id, c.complaint_id, p.sku, cp.capa_type, cp.title,
                   cp.department, cp.priority, cp.status, cp.due_date, u.full_name as assigned_to
            FROM capa_tasks cp
            JOIN complaints c ON cp.complaint_id = c.id
            JOIN products p ON c.product_id = p.id
            LEFT JOIN users u ON cp.assigned_to_user_id = u.id
            ORDER BY cp.created_at DESC
            """,
            conn
        )

        products_df = pd.read_sql_query(
            """
            SELECT p.sku, p.name, p.category, p.status, p.revision,
                   (SELECT COUNT(*) FROM documents WHERE product_id = p.id) as document_count,
                   (SELECT COUNT(*) FROM complaints WHERE product_id = p.id) as complaint_count
            FROM products p
            ORDER BY p.name ASC
            """,
            conn
        )

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        complaints_df.to_excel(writer, sheet_name='Complaints', index=False)
        evaluations_df.to_excel(writer, sheet_name='Evaluations', index=False)
        capas_df.to_excel(writer, sheet_name='CAPA Tasks', index=False)
        products_df.to_excel(writer, sheet_name='Products', index=False)
        
        # Apply enterprise styling
        for sheetname in writer.sheets:
            style_excel_worksheet(writer.sheets[sheetname])

    output.seek(0)
    return output.getvalue()

def generate_complaints_csv() -> str:
    """Exports all complaints as CSV string."""
    with get_db() as conn:
        df = pd.read_sql_query(
            """
            SELECT c.complaint_id, c.date_received, c.reporter_type, cust.name as customer_name,
                   p.sku, p.name as product_name, c.category, c.severity, c.current_status,
                   c.quantity_affected, c.description
            FROM complaints c
            JOIN products p ON c.product_id = p.id
            LEFT JOIN customers cust ON c.customer_id = cust.id
            ORDER BY c.created_at DESC
            """,
            conn
        )
    return df.to_csv(index=False)
